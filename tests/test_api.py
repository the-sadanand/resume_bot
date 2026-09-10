"""
FastAPI Integration Tests
Tests the API endpoints using TestClient
"""
import os
import pytest
from fastapi.testclient import TestClient

# Set test env before importing app
os.environ["DATABASE_URL"] = "sqlite:///./data/test_api.db"
os.environ["LLM_PROVIDER"] = "none"
os.environ["REDIS_ENABLED"] = "false"

from app.main import app

client = TestClient(app)


def test_health_endpoint():
    response = client.get("/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "ok"
    assert "version" in data
    assert "app_name" in data


def test_root_endpoint():
    response = client.get("/")
    assert response.status_code == 200
    data = response.json()
    assert "app" in data
    assert "docs" in data


def test_create_job():
    response = client.post("/api/v1/jobs", json={
        "title": "Senior Python Engineer",
        "description": """
        Requirements:
        - 5+ years Python experience
        - FastAPI, Docker, PostgreSQL required
        - AWS cloud experience
        - Kubernetes preferred
        
        Education: Bachelor's in Computer Science
        """
    })
    assert response.status_code == 200
    data = response.json()
    assert "job_id" in data
    assert data["title"] == "Senior Python Engineer"
    assert "parsed" in data
    return data["job_id"]


def test_create_job_and_retrieve():
    # Create
    create_resp = client.post("/api/v1/jobs", json={
        "title": "Test Engineer Role",
        "description": "Python developer needed. Docker, FastAPI, PostgreSQL required. 3+ years experience."
    })
    assert create_resp.status_code == 200
    job_id = create_resp.json()["job_id"]

    # Retrieve
    get_resp = client.get(f"/api/v1/jobs/{job_id}")
    assert get_resp.status_code == 200
    assert get_resp.json()["job_id"] == job_id


def test_get_nonexistent_job():
    response = client.get("/api/v1/jobs/nonexistent_job_id")
    assert response.status_code == 404


def test_upload_invalid_file_type():
    response = client.post(
        "/api/v1/resumes/upload",
        files={"file": ("resume.exe", b"malicious content", "application/octet-stream")},
    )
    assert response.status_code == 400


def test_upload_docx_resume(tmp_path):
    """Test uploading a DOCX resume created with python-docx."""
    try:
        from docx import Document as DocxDocument
        doc = DocxDocument()
        doc.add_heading("Test Candidate", 0)
        doc.add_paragraph("test@example.com")
        doc.add_heading("Skills", 2)
        doc.add_paragraph("Python, FastAPI, Docker, PostgreSQL")
        doc.add_heading("Experience", 2)
        doc.add_paragraph("Backend Engineer — TestCorp (2020–2024)")
        doc.add_paragraph("3 years Python experience")
        
        docx_path = tmp_path / "test_resume.docx"
        doc.save(str(docx_path))

        with open(docx_path, "rb") as f:
            response = client.post(
                "/api/v1/resumes/upload",
                files={"file": ("test_resume.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )

        assert response.status_code == 200
        data = response.json()
        assert "resume_id" in data
        assert data["status"] == "parsed"
        return data["resume_id"]
    except ImportError:
        pytest.skip("python-docx not installed")


def test_screen_single_resume(tmp_path):
    """End-to-end test: create job, upload resume, screen."""
    # Create job
    job_resp = client.post("/api/v1/jobs", json={
        "title": "Python Developer",
        "description": "Python, FastAPI, Docker required. 3+ years experience."
    })
    assert job_resp.status_code == 200
    job_id = job_resp.json()["job_id"]

    # Upload resume
    try:
        from docx import Document as DocxDoc
        doc = DocxDoc()
        doc.add_heading("Jane Developer", 0)
        doc.add_paragraph("jane@example.com | +1-555-0111")
        doc.add_heading("Skills", 2)
        doc.add_paragraph("Python, FastAPI, Docker, PostgreSQL, Redis, Git, Linux")
        doc.add_heading("Experience", 2)
        doc.add_paragraph("Senior Python Developer — TechCo (2020–2024)")
        doc.add_paragraph("4 years Python, FastAPI development, Docker deployment")
        doc.add_heading("Projects", 2)
        doc.add_paragraph("API Gateway — Python, FastAPI, Docker")
        doc.add_heading("Education", 2)
        doc.add_paragraph("Bachelor of Science Computer Science 2020")

        docx_path = tmp_path / "jane.docx"
        doc.save(str(docx_path))

        with open(docx_path, "rb") as f:
            resume_resp = client.post(
                "/api/v1/resumes/upload",
                files={"file": ("jane.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        assert resume_resp.status_code == 200
        resume_id = resume_resp.json()["resume_id"]

        # Screen
        screen_resp = client.post("/api/v1/screen", json={
            "job_id": job_id,
            "resume_ids": [resume_id]
        })
        assert screen_resp.status_code == 200
        result = screen_resp.json()
        assert "overall_score" in result
        assert 0 <= result["overall_score"] <= 100
        assert "section_scores" in result
        assert "skill_match" in result
        assert "strengths" in result
        assert "gaps" in result
        assert result["recommendation"] in ["Strong Match", "Good Match", "Moderate Match", "Weak Match"]

        # Get report
        analysis_id = result["analysis_id"]
        report_resp = client.get(f"/api/v1/results/{analysis_id}/report")
        assert report_resp.status_code == 200
        report_text = report_resp.json()["report"]
        assert "RESUME SCREENING RESULT" in report_text
        assert "Overall Score" in report_text

    except ImportError:
        pytest.skip("python-docx not installed")


def test_batch_screening_requires_min_two():
    response = client.post("/api/v1/screen/batch", json={
        "job_id": "any_job",
        "resume_ids": ["res_001"]
    })
    assert response.status_code == 400


def test_screen_nonexistent_job():
    response = client.post("/api/v1/screen", json={
        "job_id": "nonexistent_job",
        "resume_ids": ["res_001"]
    })
    assert response.status_code in [400, 404]


def test_get_nonexistent_analysis():
    response = client.get("/api/v1/results/nonexistent_analysis_id")
    assert response.status_code == 404


def test_list_jobs():
    response = client.get("/api/v1/jobs")
    assert response.status_code == 200
    assert isinstance(response.json(), list)
