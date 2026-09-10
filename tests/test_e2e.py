"""
End-to-End Integration Tests
Tests the complete screening pipeline
"""
import os
import pytest
import tempfile

os.environ["DATABASE_URL"] = "sqlite:///./data/test_e2e.db"
os.environ["LLM_PROVIDER"] = "none"
os.environ["REDIS_ENABLED"] = "false"


def create_docx_resume(path, name, email, skills, years_exp, has_education=True, has_projects=True):
    """Helper to create a test DOCX resume."""
    from docx import Document
    doc = Document()
    doc.add_heading(name, 0)
    doc.add_paragraph(f"{email} | +1-555-0100")
    doc.add_heading("Summary", 2)
    doc.add_paragraph(f"Experienced developer with {years_exp} years.")
    doc.add_heading("Skills", 2)
    doc.add_paragraph(", ".join(skills))
    doc.add_heading("Experience", 2)
    doc.add_paragraph(f"Senior Engineer — TestCorp (2019–2024)")
    doc.add_paragraph(f"{years_exp} years experience. Technologies: {', '.join(skills[:4])}")
    if has_projects:
        doc.add_heading("Projects", 2)
        doc.add_paragraph(f"Project Alpha — Technologies: {', '.join(skills[:3])}")
        doc.add_paragraph("Description: Scaled system to handle millions of requests.")
    if has_education:
        doc.add_heading("Education", 2)
        doc.add_paragraph("Bachelor of Science in Computer Science — University 2019")
    doc.save(path)
    return path


def test_end_to_end_single_resume(tmp_path):
    """Complete pipeline: JD → Upload → Screen → Score → Report."""
    try:
        from fastapi.testclient import TestClient
        from app.main import app
        client = TestClient(app)

        # Step 1: Create JD
        jd_resp = client.post("/api/v1/jobs", json={
            "title": "Python Backend Engineer",
            "description": (
                "Python Backend Engineer\n\n"
                "Requirements:\n"
                "- 4+ years Python experience\n"
                "- FastAPI, Django\n"
                "- Docker, PostgreSQL, Redis\n"
                "- AWS cloud\n"
                "- REST APIs\n\n"
                "Education: Bachelor's Computer Science"
            )
        })
        assert jd_resp.status_code == 200
        job_id = jd_resp.json()["job_id"]

        # Step 2: Create and upload resume
        resume_path = str(tmp_path / "test_e2e.docx")
        create_docx_resume(
            resume_path,
            name="Emma Rodriguez",
            email="emma@example.com",
            skills=["python", "fastapi", "docker", "postgresql", "redis", "aws", "git", "linux"],
            years_exp=5,
        )

        with open(resume_path, "rb") as f:
            upload_resp = client.post(
                "/api/v1/resumes/upload",
                files={"file": ("emma.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
        assert upload_resp.status_code == 200
        resume_id = upload_resp.json()["resume_id"]

        # Step 3: Screen
        screen_resp = client.post("/api/v1/screen", json={
            "job_id": job_id,
            "resume_ids": [resume_id],
        })
        assert screen_resp.status_code == 200
        result = screen_resp.json()

        # Verify result structure
        assert "analysis_id" in result
        assert "overall_score" in result
        assert "section_scores" in result
        assert "skill_match" in result
        assert "strengths" in result
        assert "gaps" in result

        overall = result["overall_score"]
        assert 0 <= overall <= 100

        # Emma has matching skills — should be at least Moderate
        assert result["recommendation"] in ["Strong Match", "Good Match", "Moderate Match"]

        # Verify section scores exist
        sections = result["section_scores"]
        assert "skills" in sections
        assert "experience" in sections
        assert "projects" in sections
        assert "education" in sections

        # Step 4: Get full report
        analysis_id = result["analysis_id"]
        report_resp = client.get(f"/api/v1/results/{analysis_id}/report")
        assert report_resp.status_code == 200
        report = report_resp.json()["report"]
        assert "Emma Rodriguez" in report
        assert "RESUME SCREENING RESULT" in report
        assert "Overall Score" in report
        assert "MATCHED SKILLS" in report

        print(f"\n✅ E2E Single Resume Test PASSED")
        print(f"   Candidate: Emma Rodriguez")
        print(f"   Score: {overall:.1f}/100")
        print(f"   Recommendation: {result['recommendation']}")

    except ImportError as e:
        pytest.skip(f"Required package not installed: {e}")


def test_end_to_end_batch_screening(tmp_path):
    """Batch pipeline: JD → 3 Resumes → Rank → Charts."""
    try:
        from fastapi.testclient import TestClient
        from app.main import app
        client = TestClient(app)

        # Create JD
        jd_resp = client.post("/api/v1/jobs", json={
            "title": "Senior Python Backend Engineer",
            "description": (
                "Senior Python Backend Engineer\n\n"
                "Required: Python, FastAPI, Docker, PostgreSQL, Redis, AWS, CI/CD\n"
                "Preferred: Kubernetes, Elasticsearch, GraphQL\n"
                "5+ years experience\n"
                "Education: Bachelor's Computer Science"
            )
        })
        assert jd_resp.status_code == 200
        job_id = jd_resp.json()["job_id"]

        # Create 3 resumes with different skill levels
        candidates = [
            {
                "name": "Alex Strong",
                "skills": ["python", "fastapi", "docker", "postgresql", "redis", "aws", "kubernetes", "elasticsearch"],
                "years": 7,
            },
            {
                "name": "Morgan Medium",
                "skills": ["python", "django", "docker", "postgresql", "git", "linux"],
                "years": 3,
            },
            {
                "name": "Taylor Weak",
                "skills": ["javascript", "react", "nodejs", "mongodb"],
                "years": 2,
            },
        ]

        resume_ids = []
        for c in candidates:
            path = str(tmp_path / f"{c['name'].lower().replace(' ', '_')}.docx")
            create_docx_resume(path, c["name"], f"{c['name'].lower().replace(' ', '')}@example.com",
                               c["skills"], c["years"])
            with open(path, "rb") as f:
                r = client.post(
                    "/api/v1/resumes/upload",
                    files={"file": (f"{c['name']}.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                )
            assert r.status_code == 200
            resume_ids.append(r.json()["resume_id"])

        # Batch screen
        batch_resp = client.post("/api/v1/screen/batch", json={
            "job_id": job_id,
            "resume_ids": resume_ids,
        })
        assert batch_resp.status_code == 200
        result = batch_resp.json()

        assert result["total_candidates"] == 3
        assert len(result["ranked_candidates"]) == 3
        assert len(result["comparison_table"]) == 3

        # Verify ranking is sorted
        scores = [rc["overall_score"] for rc in result["ranked_candidates"]]
        assert scores == sorted(scores, reverse=True), f"Scores not sorted: {scores}"

        # Verify charts generated
        chart_paths = result.get("chart_paths", [])
        for path in chart_paths:
            assert os.path.exists(path), f"Chart file not found: {path}"

        print(f"\n✅ E2E Batch Screening Test PASSED")
        print(f"   Candidates ranked:")
        for rc in result["ranked_candidates"]:
            print(f"   {rc['rank']}. {rc['candidate_name']} — {rc['overall_score']:.1f}/100 ({rc['recommendation']})")

        if chart_paths:
            print(f"   Charts generated: {len(chart_paths)}")

    except ImportError as e:
        pytest.skip(f"Required package not installed: {e}")
