"""Telegram-only resume screening showcase with visual recruiter reports."""
import io
import re
from typing import Optional
import httpx
from fastapi import APIRouter, HTTPException, Request
from app.core.config import get_settings
from app.core.logging import get_logger

router = APIRouter(prefix="/integrations/telegram", tags=["Telegram"])
settings = get_settings(); logger = get_logger(__name__)
SKILLS = ["python","java","javascript","typescript","c++","c#","go","rust","fastapi","django","flask","react","angular","node.js","nodejs","docker","kubernetes","aws","azure","gcp","git","linux","mysql","postgresql","mongodb","redis","sql","nosql","tensorflow","pytorch","scikit-learn","pandas","numpy","keras","machine learning","deep learning","nlp","computer vision","data science","rest api","graphql","microservices","celery","kafka","spark"]
COURSES = {
 "python":("Python","https://www.python.org/about/gettingstarted/"), "fastapi":("FastAPI","https://fastapi.tiangolo.com/tutorial/"), "docker":("Docker","https://docs.docker.com/get-started/"), "git":("Git","https://git-scm.com/docs/gittutorial"), "aws":("AWS","https://aws.amazon.com/training/"), "sql":("SQL","https://www.w3schools.com/sql/"), "pytorch":("PyTorch","https://pytorch.org/tutorials/"), "tensorflow":("TensorFlow","https://www.tensorflow.org/learn"), "scikit-learn":("Scikit-learn","https://scikit-learn.org/stable/getting_started.html"), "pandas":("Pandas","https://pandas.pydata.org/docs/getting_started/intro_tutorials/"), "numpy":("NumPy","https://numpy.org/learn/"), "machine learning":("Machine Learning","https://scikit-learn.org/stable/getting_started.html"), "react":("React","https://react.dev/learn"), "java":("Java","https://dev.java/learn/"), "javascript":("JavaScript","https://developer.mozilla.org/en-US/docs/Web/JavaScript/Guide"), "kubernetes":("Kubernetes","https://kubernetes.io/docs/tutorials/"), "linux":("Linux","https://ubuntu.com/tutorials/command-line-for-beginners")}
SESSIONS: dict[str, dict] = {}

def session_for(uid): return SESSIONS.setdefault(uid,{"state":"idle","jd":"","resumes":[]})
def normalize(text): return re.sub(r"\s+"," ",text.lower()).strip()
def contains_skill(text,skill):
 text,skill=normalize(text),skill.lower()
 if " " in skill or "." in skill or "+" in skill or "#" in skill: return skill in text
 return bool(re.search(r"(?<![a-z0-9])"+re.escape(skill)+r"(?![a-z0-9])",text))
def extract_skills(text): return [s for s in SKILLS if contains_skill(text,s)]
def extract_years(text): return max((float(x) for x in re.findall(r"(\d+(?:\.\d+)?)\s*\+?\s*years?",text.lower())),default=0.0)
def extract_candidate_name(text,filename):
 for line in [x.strip() for x in text.splitlines() if x.strip()][:10]:
  low=line.lower()
  if 2<=len(line.split())<=5 and len(line)<60 and not any(x in low for x in ["resume","curriculum","email","phone","linkedin","github"]) and not any(c.isdigit() for c in line): return line
 return re.sub(r"[_-]+"," ",filename.rsplit(".",1)[0]).strip() or "Candidate"
def extract_resume_text(data,filename):
 suffix=filename.lower().rsplit(".",1)[-1]
 if suffix=="pdf":
  import fitz
  with fitz.open(stream=data,filetype="pdf") as doc: return "\n".join(p.get_text() for p in doc)
 if suffix=="docx":
  from docx import Document
  doc=Document(io.BytesIO(data)); parts=[p.text for p in doc.paragraphs if p.text.strip()]
  for table in doc.tables:
   for row in table.rows:
    for cell in row.cells:
     if cell.text.strip(): parts.append(cell.text.strip())
  return "\n".join(parts)
 raise ValueError("Only PDF and DOCX files are supported")
def score_resume(jd,resume):
 jd_skills,resume_skills=extract_skills(jd),extract_skills(resume); matched=[s for s in jd_skills if s in resume_skills]; missing=[s for s in jd_skills if s not in resume_skills]
 skills=len(matched)/len(jd_skills)*100 if jd_skills else 0.0; required,resume_y=extract_years(jd),extract_years(resume)
 experience=min(100.0,resume_y/required*100) if required else (100.0 if resume_y else 60.0)
 projects=min(100.0,40.0+sum(x in normalize(resume) for x in ["project","developed","built","implemented"])*15.0)
 education=100.0 if any(x in normalize(resume) for x in ["b.tech","btech","bachelor","b.e","m.tech","mtech","master","degree"]) else 40.0
 completeness=sum([bool(extract_candidate_name(resume,"candidate.pdf")),"@" in resume,bool(re.search(r"\b(?:experience|work)\b",resume,re.I)),bool(re.search(r"\b(?:education|b\.tech|bachelor|degree)\b",resume,re.I)),bool(resume_skills)])/5*100
 overall=skills*.35+experience*.20+projects*.15+education*.10+skills*.15+completeness*.05
 rec="STRONG MATCH" if overall>=85 else "GOOD MATCH" if overall>=70 else "MODERATE MATCH" if overall>=55 else "WEAK MATCH"
 strengths=[]; gaps=[]
 if matched: strengths.append("Strongest match: "+", ".join(matched[:5]))
 if experience>=80: strengths.append("Experience level matches the requirement")
 if projects>=70: strengths.append("Relevant project evidence found")
 if missing: gaps.append("Missing: "+", ".join(missing[:6]))
 if required and resume_y<required: gaps.append(f"Experience: {resume_y:g} years found vs {required:g} required")
 if not gaps: gaps.append("No major gaps detected")
 return {"overall":overall,"skills":skills,"experience":experience,"projects":projects,"education":education,"jd_match":skills,"completeness":completeness,"matched":matched,"missing":missing,"strengths":strengths,"gaps":gaps,"recommendation":rec}
def bar(value,width=12):
 filled=round(max(0,min(100,value))/100*width); return "█"*filled+"░"*(width-filled)
def courses(missing):
 out=[]
 for skill in missing:
  if skill in COURSES and COURSES[skill] not in out: out.append(COURSES[skill])
  if len(out)>=4: break
 return out
def format_result(candidate,r):
 lines=["━━━━━━━━━━━━━━━━━━━━","🤖 RESUME INSIGHT","━━━━━━━━━━━━━━━━━━━━","",f"👤 {candidate}",f"🎯 OVERALL MATCH: {r['overall']:.1f}/100",f"📌 {r['recommendation']}","","📊 CATEGORY STRENGTH","",f"Skills       {bar(r['skills'])} {r['skills']:.0f}%",f"Experience   {bar(r['experience'])} {r['experience']:.0f}%",f"Projects     {bar(r['projects'])} {r['projects']:.0f}%",f"Education    {bar(r['education'])} {r['education']:.0f}%",f"JD Match     {bar(r['jd_match'])} {r['jd_match']:.0f}%",f"Completeness {bar(r['completeness'])} {r['completeness']:.0f}%","","💻 SKILL STRENGTH",""]
 for s in r["matched"][:8]: lines.append(f"✅ {s:<18} {bar(100)}  STRONG")
 for s in r["missing"][:8]: lines.append(f"❌ {s:<18} {bar(0)}  NEEDS WORK")
 lines += ["","💪 STRENGTHS"]+[f"• {x}" for x in r["strengths"][:3]]+["","⚠️ GAPS"]+[f"• {x}" for x in r["gaps"][:3]]
 recs=courses(r["missing"])
 if recs:
  lines += ["","🎓 RECOMMENDED LEARNING"]+[f"• {n}: {u}" for n,u in recs]
 lines += ["","📌 FINAL VERDICT",r["recommendation"],"━━━━━━━━━━━━━━━━━━━━"]
 return "\n".join(lines)
def comparison_chart(results):
 from PIL import Image,ImageDraw,ImageFont
 W,H=900,120+95*len(results); image=Image.new("RGB",(W,H),"white"); draw=ImageDraw.Draw(image); font=ImageFont.load_default(size=24); small=ImageFont.load_default(size=18)
 draw.text((35,25),"RESUME MATCH COMPARISON",fill="black",font=font)
 for i,(name,r) in enumerate(results):
  y=95+i*95; draw.text((35,y),f"#{i+1} {name[:28]}",fill="black",font=small); x,bw=300,500; draw.rectangle((x,y,x+bw,y+30),outline="black",width=2); fill=int(bw*r["overall"]/100)
  if fill: draw.rectangle((x,y,x+fill,y+30),fill="black")
  draw.text((815,y+5),f"{r['overall']:.0f}%",fill="black",font=small)
 out=io.BytesIO(); image.save(out,format="PNG"); return out.getvalue()
async def send_message(chat_id,text):
 if not settings.telegram_bot_token: return False
 try:
  async with httpx.AsyncClient(timeout=20) as client:
   r=await client.post(f"https://api.telegram.org/bot{settings.telegram_bot_token}/sendMessage",json={"chat_id":chat_id,"text":text,"disable_web_page_preview":True}); return r.status_code==200
 except Exception as exc: logger.error(f"Telegram send failed: {type(exc).__name__}"); return False
async def send_photo(chat_id,data,caption=""):
 if not settings.telegram_bot_token: return False
 try:
  async with httpx.AsyncClient(timeout=30) as client:
   r=await client.post(f"https://api.telegram.org/bot{settings.telegram_bot_token}/sendPhoto",data={"chat_id":chat_id,"caption":caption},files={"photo":("comparison.png",data,"image/png")}); return r.status_code==200
 except Exception as exc: logger.error(f"Telegram photo failed: {type(exc).__name__}"); return False
async def download_file(file_id):
 if not settings.telegram_bot_token: return None
 try:
  async with httpx.AsyncClient(timeout=30) as client:
   meta=await client.get(f"https://api.telegram.org/bot{settings.telegram_bot_token}/getFile",params={"file_id":file_id}); meta.raise_for_status(); path=meta.json()["result"]["file_path"]
   r=await client.get(f"https://api.telegram.org/file/bot{settings.telegram_bot_token}/{path}"); r.raise_for_status(); return r.content
 except Exception as exc: logger.error(f"Telegram download failed: {type(exc).__name__}"); return None
@router.post("/webhook")
async def telegram_webhook(request:Request):
 try: update=await request.json()
 except Exception as exc: raise HTTPException(status_code=400,detail="Invalid JSON") from exc
 message=update.get("message") or update.get("edited_message")
 if not message: return {"ok":True}
 chat_id,user_id=str(message["chat"]["id"]),str(message["from"]["id"]); text=message.get("text","").strip(); document=message.get("document"); session=session_for(user_id)
 if text.startswith("/start") or text.startswith("/help"):
  await send_message(chat_id,"👋 Resume Screening Bot\n\n/analyze - start screening\n/screen - rank uploaded resumes\n/reset - start over\n\n📊 Report includes visual comparison, filled skill bars, gaps and course recommendations."); return {"ok":True}
 if text.startswith("/reset"):
  session.update({"state":"idle","jd":"","resumes":[]}); await send_message(chat_id,"✅ Reset complete. Use /analyze to start."); return {"ok":True}
 if text.startswith("/analyze"):
  session.update({"state":"awaiting_jd","jd":"","resumes":[]}); await send_message(chat_id,"📋 Send the Job Description text now."); return {"ok":True}
 if session["state"]=="awaiting_jd" and text and not text.startswith("/"):
  session["jd"],session["state"]=text,"awaiting_resumes"; skills=extract_skills(text); await send_message(chat_id,"✅ Job Description received\n\n🎯 Detected skills:\n"+(", ".join(skills[:15]) if skills else "None")+"\n\n📎 Upload PDF/DOCX resumes, then send /screen."); return {"ok":True}
 if session["state"]=="awaiting_resumes":
  if document:
   filename=document.get("file_name","resume.pdf")
   if not filename.lower().endswith((".pdf",".docx")): await send_message(chat_id,"❌ Only PDF and DOCX files are supported."); return {"ok":True}
   data=await download_file(document["file_id"])
   if not data: await send_message(chat_id,"❌ Could not download the resume. Try again."); return {"ok":True}
   try:
    resume_text=extract_resume_text(data,filename)
    if not resume_text.strip(): raise ValueError("No text found")
    candidate=extract_candidate_name(resume_text,filename); session["resumes"].append({"name":candidate,"text":resume_text}); await send_message(chat_id,f"✅ Resume received: {candidate}\n💻 Skills found: {len(extract_skills(resume_text))}\n\nSend another resume or /screen.")
   except Exception as exc: logger.error(f"Resume parsing failed: {type(exc).__name__}"); await send_message(chat_id,"❌ Could not read this resume. Try another PDF/DOCX.")
   return {"ok":True}
  if text.startswith("/screen"):
   if not session["resumes"]: await send_message(chat_id,"❌ Upload at least one resume first."); return {"ok":True}
   await send_message(chat_id,f"⏳ Screening {len(session['resumes'])} resume(s)...")
   results=[(r["name"],score_resume(session["jd"],r["text"])) for r in session["resumes"]]; results.sort(key=lambda x:x[1]["overall"],reverse=True)
   if len(results)>1:
    try: await send_photo(chat_id,comparison_chart(results),"📊 Candidate comparison — ranked by overall match")
    except Exception as exc: logger.error(f"Chart generation failed: {type(exc).__name__}")
   for i,(candidate,result) in enumerate(results,1): await send_message(chat_id,(f"🏆 RANK #{i}\n\n" if len(results)>1 else "")+format_result(candidate,result))
   session.update({"state":"idle","jd":"","resumes":[]}); return {"ok":True}
 if session["state"]=="idle": await send_message(chat_id,"Use /analyze to start resume screening.")
 return {"ok":True}
