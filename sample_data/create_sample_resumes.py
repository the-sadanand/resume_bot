"""
Script to generate sample resume DOCX files for testing.
Uses only fictional candidates - no real personal information.
"""
import os
import sys

# Add parent directory to path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("python-docx not installed. Run: pip install python-docx")
    sys.exit(1)


def add_heading(doc, text, level=1, bold=True):
    heading = doc.add_heading(text, level=level)
    if bold:
        for run in heading.runs:
            run.bold = True


def add_para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
    p.add_run(text)
    return p


def create_resume_alex_chen():
    """Strong Match candidate — Senior Python Engineer."""
    doc = Document()

    # Name & contact
    doc.add_heading("Alex Chen", 0)
    doc.add_paragraph("alex.chen@example.com  |  +1-555-0101  |  linkedin.com/in/alexchen-fictional")

    # Summary
    add_heading(doc, "Professional Summary", 2)
    doc.add_paragraph(
        "Seasoned Python Backend Engineer with 7 years of experience building high-performance, "
        "scalable APIs and distributed systems. Deep expertise in FastAPI, PostgreSQL, Docker, "
        "and AWS. Proven track record delivering production systems processing millions of requests daily."
    )

    # Skills
    add_heading(doc, "Technical Skills", 2)
    doc.add_paragraph("Python, FastAPI, Django, Flask, PostgreSQL, Redis, Docker, Kubernetes")
    doc.add_paragraph("AWS (EC2, RDS, S3, Lambda), GCP, Celery, Elasticsearch, GraphQL, RabbitMQ")
    doc.add_paragraph("SQLAlchemy, Alembic, pytest, GitHub Actions, Jenkins, Terraform")
    doc.add_paragraph("Git, Linux, Bash, Nginx, REST APIs, Microservices")

    # Experience
    add_heading(doc, "Work Experience", 2)

    add_heading(doc, "Senior Backend Engineer — CloudScale Technologies (2020–Present)", 3)
    doc.add_paragraph("• Architected and built a FastAPI-based microservices platform serving 5M requests/day")
    doc.add_paragraph("• Designed PostgreSQL schemas with Redis caching, reducing API latency by 60%")
    doc.add_paragraph("• Led Docker and Kubernetes migration, achieving 99.9% uptime")
    doc.add_paragraph("• Built CI/CD pipelines using GitHub Actions and Jenkins")
    doc.add_paragraph("• Mentored 4 junior engineers")

    add_heading(doc, "Backend Engineer — DataStream Inc (2018–2020)", 3)
    doc.add_paragraph("• Developed Django REST APIs for financial data processing")
    doc.add_paragraph("• Implemented Celery task queues with RabbitMQ for async processing")
    doc.add_paragraph("• Integrated Elasticsearch for full-text search across 50M+ records")

    add_heading(doc, "Junior Python Developer — TechStart Labs (2017–2018)", 3)
    doc.add_paragraph("• Built Flask APIs for e-commerce backend")
    doc.add_paragraph("• Wrote unit tests using pytest, achieving 85% code coverage")

    # Projects
    add_heading(doc, "Projects", 2)
    add_heading(doc, "AutoScaler — Open Source Kubernetes Autoscaling Tool", 3)
    doc.add_paragraph("Technologies: Python, Kubernetes, Docker, AWS, Go")
    doc.add_paragraph("Built a custom Kubernetes horizontal pod autoscaler based on business metrics.")

    add_heading(doc, "RealtimeDB — Event-Driven Data Pipeline", 3)
    doc.add_paragraph("Technologies: Python, Kafka, PostgreSQL, Redis, FastAPI")
    doc.add_paragraph("Designed event streaming pipeline processing 100K events/sec.")

    # Education
    add_heading(doc, "Education", 2)
    doc.add_paragraph("Bachelor of Science in Computer Science — State University (2017)")

    # Certifications
    add_heading(doc, "Certifications", 2)
    doc.add_paragraph("• AWS Certified Solutions Architect — Associate (2022)")
    doc.add_paragraph("• Certified Kubernetes Administrator (CKA) (2021)")

    # Achievements
    add_heading(doc, "Achievements", 2)
    doc.add_paragraph("• Speaker at PyCon 2023: 'Building High-Performance FastAPI Applications'")
    doc.add_paragraph("• Open source contributor: 500+ GitHub stars on AutoScaler project")

    return doc


def create_resume_priya_sharma():
    """Good Match candidate — Python developer with some gaps."""
    doc = Document()

    doc.add_heading("Priya Sharma", 0)
    doc.add_paragraph("priya.sharma@example.com  |  +91-98765-43210")

    add_heading(doc, "Summary", 2)
    doc.add_paragraph(
        "Python developer with 4 years of experience in backend development. "
        "Proficient in Django, FastAPI, and MySQL. Eager to grow in cloud technologies."
    )

    add_heading(doc, "Skills", 2)
    doc.add_paragraph("Python, FastAPI, Django, MySQL, MongoDB, Docker")
    doc.add_paragraph("REST APIs, Git, Linux, pytest, SQLAlchemy, HTML, CSS")
    doc.add_paragraph("AWS (basic), Redis (basic)")

    add_heading(doc, "Experience", 2)

    add_heading(doc, "Python Backend Developer — Innovatech Solutions (2022–Present)", 3)
    doc.add_paragraph("• Built FastAPI services for B2B SaaS platform")
    doc.add_paragraph("• Implemented Redis caching for API performance")
    doc.add_paragraph("• Docker containerization for development environments")
    doc.add_paragraph("• MySQL database design and query optimization")

    add_heading(doc, "Junior Python Developer — WebSoft India (2020–2022)", 3)
    doc.add_paragraph("• Developed Django REST framework APIs")
    doc.add_paragraph("• Integrated payment gateways (Razorpay, Stripe)")
    doc.add_paragraph("• Wrote unit tests and documentation")

    add_heading(doc, "Projects", 2)
    add_heading(doc, "HR Management System", 3)
    doc.add_paragraph("Technologies: Django, MySQL, Docker, REST API")
    doc.add_paragraph("Full-stack HR portal for attendance, payroll, and performance tracking.")

    add_heading(doc, "E-Commerce API", 3)
    doc.add_paragraph("Technologies: FastAPI, Redis, PostgreSQL")
    doc.add_paragraph("RESTful API for multi-vendor e-commerce platform.")

    add_heading(doc, "Education", 2)
    doc.add_paragraph("Bachelor of Engineering in Computer Science — VIT University (2020)")

    add_heading(doc, "Certifications", 2)
    doc.add_paragraph("• AWS Cloud Practitioner (2023)")

    return doc


def create_resume_james_miller():
    """Moderate Match — different background, partial overlap."""
    doc = Document()

    doc.add_heading("James Miller", 0)
    doc.add_paragraph("james.miller@example.com  |  +1-555-0202")

    add_heading(doc, "Profile", 2)
    doc.add_paragraph(
        "Full-stack developer with 3 years of experience, primarily in JavaScript and Node.js. "
        "Transitioning to Python backend development. Familiar with REST APIs and databases."
    )

    add_heading(doc, "Technical Skills", 2)
    doc.add_paragraph("JavaScript, TypeScript, Node.js, Express.js, React, Vue.js")
    doc.add_paragraph("Python (intermediate), Flask, MongoDB, MySQL, PostgreSQL (basic)")
    doc.add_paragraph("Docker (basic), Git, REST APIs, Jest, AWS (basic)")

    add_heading(doc, "Work Experience", 2)

    add_heading(doc, "Full Stack Developer — DigitalEdge Agency (2022–Present)", 3)
    doc.add_paragraph("• Built Node.js/Express REST APIs for mobile applications")
    doc.add_paragraph("• React frontend development for 3 client projects")
    doc.add_paragraph("• MongoDB database design and optimization")
    doc.add_paragraph("• Basic Docker for local development")

    add_heading(doc, "Junior Developer — WebWorks Co (2021–2022)", 3)
    doc.add_paragraph("• JavaScript frontend and Node.js backend development")
    doc.add_paragraph("• MySQL database integration")

    add_heading(doc, "Projects", 2)
    add_heading(doc, "Task Management App", 3)
    doc.add_paragraph("Technologies: Python, Flask, PostgreSQL, React")
    doc.add_paragraph("Personal project: task management app with Python Flask backend.")

    add_heading(doc, "Social Media Dashboard", 3)
    doc.add_paragraph("Technologies: Node.js, MongoDB, React")
    doc.add_paragraph("Analytics dashboard for social media metrics.")

    add_heading(doc, "Education", 2)
    doc.add_paragraph("Bachelor of Science in Information Technology — City College (2021)")

    return doc


def create_resume_fatima_al_rashid():
    """Weak Match — mostly data science, minimal backend overlap."""
    doc = Document()

    doc.add_heading("Fatima Al-Rashid", 0)
    doc.add_paragraph("fatima.alrashid@example.com  |  +971-50-123-4567")

    add_heading(doc, "Professional Summary", 2)
    doc.add_paragraph(
        "Data Scientist with 3 years of experience in machine learning and data analysis. "
        "Proficient in Python for data science workflows. Limited backend API experience."
    )

    add_heading(doc, "Skills", 2)
    doc.add_paragraph("Python, Pandas, NumPy, Scikit-learn, TensorFlow, Keras")
    doc.add_paragraph("Jupyter Notebook, Matplotlib, Seaborn, SQL, R")
    doc.add_paragraph("Flask (basic), Git, Linux (basic)")

    add_heading(doc, "Experience", 2)

    add_heading(doc, "Data Scientist — Analytics Corp (2022–Present)", 3)
    doc.add_paragraph("• Built ML models for customer churn prediction (85% accuracy)")
    doc.add_paragraph("• EDA and visualization using Pandas, Matplotlib, Seaborn")
    doc.add_paragraph("• SQL queries for data extraction from data warehouses")
    doc.add_paragraph("• Deployed models as Flask APIs (internal use)")

    add_heading(doc, "Data Analyst — Insight Bureau (2021–2022)", 3)
    doc.add_paragraph("• Statistical analysis and reporting using Python and R")
    doc.add_paragraph("• Built dashboards in Tableau")

    add_heading(doc, "Projects", 2)
    add_heading(doc, "Customer Churn Prediction", 3)
    doc.add_paragraph("Technologies: Python, Scikit-learn, Pandas, Flask")
    doc.add_paragraph("Trained XGBoost model, deployed via Flask API.")

    add_heading(doc, "Education", 2)
    doc.add_paragraph("Master of Science in Data Science — Analytics University (2021)")
    doc.add_paragraph("Bachelor of Science in Mathematics — National University (2019)")

    return doc


def main():
    output_dir = os.path.dirname(os.path.abspath(__file__))

    resumes = [
        ("resume_alex_chen.docx", create_resume_alex_chen),
        ("resume_priya_sharma.docx", create_resume_priya_sharma),
        ("resume_james_miller.docx", create_resume_james_miller),
        ("resume_fatima_alrashid.docx", create_resume_fatima_al_rashid),
    ]

    for filename, creator in resumes:
        path = os.path.join(output_dir, filename)
        doc = creator()
        doc.save(path)
        print(f"Created: {path}")

    print(f"\nAll {len(resumes)} sample resumes created in: {output_dir}")
    print("\nExpected screening outcomes against Senior Python Backend Engineer JD:")
    print("  Alex Chen      → Strong Match  (~85-90)")
    print("  Priya Sharma   → Good Match    (~70-80)")
    print("  James Miller   → Moderate Match (~55-65)")
    print("  Fatima Al-Rashid → Weak Match  (~35-50)")


if __name__ == "__main__":
    main()
